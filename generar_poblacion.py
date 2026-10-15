"""Genera poblacion_municipios.csv a partir de las proyecciones del DANE.

El dataset de calidad del agua no trae ninguna variable de poblacion, por eso
se usa una fuente externa: DANE, Serie municipal de poblacion por area,
periodo 2018-2042 (CNPV 2018, publicada el 8 de agosto de 2025).
Archivo original sin modificar: fuentes/PPED-AreaMun-2018-2042_VP.xlsx

Se usa openpyxl (Python puro) en lugar de pandas, igual que app.py usa el
modulo csv, para no depender de librerias compiladas.
"""
import csv

from openpyxl import load_workbook

ARCHIVO = "fuentes/PPED-AreaMun-2018-2042_VP.xlsx"
HOJA = "PobMunicipalxÁrea"
SALIDA = "poblacion_municipios.csv"
ANIO_INICIO, ANIO_FIN = 2019, 2025

# Nombre como aparece en nuestro dataset -> codigo DIVIPOLA del DANE.
# Se cruza por codigo y no por nombre porque el DANE escribe con tildes
# ("Cajicá", "Bogotá, D.C.", "San José de Cúcuta") y nuestro dataset no
# (y a veces trae "CHI@").
CODIGOS = {
    # Sabana (Regional)
    "BOGOTA D.C.": "11001",
    "CAJICA": "25126",
    "CHIA": "25175",
    "FUNZA": "25286",
    "MADRID": "25430",
    "MOSQUERA": "25473",
    "SOACHA": "25754",
    "SOPO": "25758",
    "TOCANCIPA": "25817",
    # Resto del pais (Nacional)
    "MEDELLIN": "05001",
    "BELLO": "05088",
    "BARRANQUILLA": "08001",
    "CARTAGENA": "13001",
    "TUNJA": "15001",
    "MANIZALES": "17001",
    "POPAYAN": "19001",
    "VALLEDUPAR": "20001",
    "MONTERIA": "23001",
    "NEIVA": "41001",
    "SANTA MARTA": "47001",
    "VILLAVICENCIO": "50001",
    "PASTO": "52001",
    "CUCUTA": "54001",
    "ARMENIA": "63001",
    "PEREIRA": "66001",
    "BUCARAMANGA": "68001",
    "SINCELEJO": "70001",
    "IBAGUE": "73001",
    "CALI": "76001",
}
NOMBRE_POR_CODIGO = {codigo: nombre for nombre, codigo in CODIGOS.items()}

hoja = load_workbook(ARCHIVO, read_only=True)[HOJA]

# Los encabezados estan en la fila 8 de Excel; los datos empiezan despues.
filas = hoja.iter_rows(min_row=8, values_only=True)
encabezados = next(filas)
col = {nombre: i for i, nombre in enumerate(encabezados)}

resultado = []
for fila in filas:
    codigo = fila[col["MPIO"]]  # el codigo viene como texto: "05001"
    anio = fila[col["AÑO"]]
    if fila[col["ÁREA GEOGRÁFICA"]] != "Total":
        continue  # solo el total del municipio, no cabecera ni rural
    if anio is None or not (ANIO_INICIO <= anio <= ANIO_FIN):
        continue
    if codigo not in NOMBRE_POR_CODIGO:
        continue
    resultado.append({
        "codigo_divipola": codigo,
        "municipio": NOMBRE_POR_CODIGO[codigo],
        "departamento": fila[col["DPNOM"]],
        "anio": int(anio),
        "poblacion": int(fila[col["TOTAL"]]),
    })

resultado.sort(key=lambda r: (r["municipio"], r["anio"]))

# Verificacion: 29 municipios x 7 anios = 203 filas.
esperado = len(CODIGOS) * (ANIO_FIN - ANIO_INICIO + 1)
print(f"Filas generadas: {len(resultado)} (esperado {esperado})")
assert len(resultado) == esperado, "Falta o sobra algun municipio/anio"

with open(SALIDA, "w", newline="", encoding="utf-8") as f:
    escritor = csv.DictWriter(f, fieldnames=list(resultado[0].keys()))
    escritor.writeheader()
    escritor.writerows(resultado)
print(f"Archivo guardado: {SALIDA}")
